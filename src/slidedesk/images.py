"""Render slide images directly from .pptx files using render-pptx.

Every rendered slide is saved as a PNG in the on-disk cache, keyed by the deck's
path, modification time and size, the slide index and the layout mode. Rendering
happens one slide at a time under a global lock, so the scanner, the embedding
worker and the on-demand `BackgroundRenderer` interleave instead of waiting for
whole decks, and each PNG is available as soon as its slide is rendered.
"""
from __future__ import annotations

import hashlib
import logging
import tempfile
import threading
from collections import OrderedDict
from pathlib import Path
from typing import Optional

from PIL import Image
from pptx import Presentation

logger = logging.getLogger(__name__)

# render-pptx renders one slide at a time.
_render_lock = threading.Lock()
# Decks whose rendering raised, keyed by _source_key, so they are not retried until edited.
_failures: dict = {}


class RenderError(RuntimeError):
    pass


def _source_key(pptx_path: Path, strip: bool) -> tuple:
    stat = pptx_path.stat()
    return (str(pptx_path), stat.st_mtime_ns, stat.st_size, strip)


def _cache_file(source_key: tuple, index: int, cache_dir: Path) -> Path:
    path, mtime_ns, size, strip = source_key
    raw = f"{path}|{mtime_ns}|{size}|{index}|{int(strip)}"
    return cache_dir / f"{hashlib.sha1(raw.encode('utf-8')).hexdigest()}.png"


def cached_path(pptx_path: Path, index: int, cache_dir: Path, strip: bool = False) -> Optional[Path]:
    """Return the cached PNG of slide `index` (0-based), or None if it is not rendered yet."""
    candidate = _cache_file(_source_key(pptx_path, strip), index, cache_dir)
    return candidate if candidate.exists() else None


def render_failure(pptx_path: Path, strip: bool = False) -> Optional[str]:
    """Return why rendering the current version of `pptx_path` failed, if it did."""
    return _failures.get(_source_key(pptx_path, strip))


def _open_presentation(pptx_path: Path, strip: bool, tmp: str):
    from . import layout_strip, pptx_tools

    if strip:
        stripped = Path(tmp) / pptx_path.name
        layout_strip.make_strip_layout_copy(pptx_path, stripped)
        return Presentation(str(stripped))
    prs = Presentation(str(pptx_path))
    for slide in prs.slides:
        pptx_tools.set_slide_hidden(slide, False)
    return prs


def _render_slide(prs, index: int, strip: bool) -> Image.Image:
    from render_pptx import render_slides

    image = render_slides(prs, [index])[0]
    return image.convert("L").convert("RGB") if strip else image


def render_deck(pptx_path: Path, cache_dir: Path, strip: bool = False,
                stop_event: Optional[threading.Event] = None) -> None:
    """Render and cache every slide of `pptx_path` that is not cached yet.

    Every slide is rendered as if visible, including hidden slides. With `strip`,
    the grayscale strip-layout rendering is produced. Raises RenderError if the
    deck cannot be rendered; the failure is remembered until the file changes.
    """
    key = _source_key(pptx_path, strip)
    if key in _failures:
        raise RenderError(f"Rendering {pptx_path.name} failed earlier: {_failures[key]}")
    cache_dir.mkdir(parents=True, exist_ok=True)
    try:
        with tempfile.TemporaryDirectory() as tmp:
            prs = _open_presentation(pptx_path, strip, tmp)
            for index in range(len(prs.slides)):
                if stop_event is not None and stop_event.is_set():
                    return
                target = _cache_file(key, index, cache_dir)
                with _render_lock:
                    if target.exists():  # rendered meanwhile by another thread
                        continue
                    image = _render_slide(prs, index, strip)
                    tmp_target = target.with_suffix(".tmp")
                    image.save(tmp_target, format="PNG")
                    tmp_target.replace(target)
    except Exception as exc:
        logger.exception("Rendering %s failed", pptx_path)
        _failures[key] = f"{type(exc).__name__}: {exc}"
        raise RenderError(f"Rendering {pptx_path.name} failed: {exc}") from exc


def render_slide_image(
    pptx_path: Path, index: int, cache_dir: Path, strip: bool = False
) -> Image.Image:
    """Return slide `index` (0-based) of `pptx_path` as a PIL Image, rendering the
    deck first if the slide is not cached."""
    cached = cached_path(pptx_path, index, cache_dir, strip)
    if cached is None:
        render_deck(pptx_path, cache_dir, strip)
        cached = cached_path(pptx_path, index, cache_dir, strip)
    if cached is None:
        raise ValueError(f"Slide {index} not found in {pptx_path}")
    with Image.open(cached) as image:
        image.load()
        return image.copy()


class BackgroundRenderer:
    """Renders requested decks on one daemon thread, most recently requested first."""

    def __init__(self, cache_dir: Path):
        self.cache_dir = cache_dir
        self._pending: OrderedDict = OrderedDict()
        self._cond = threading.Condition()
        self._stop = threading.Event()
        self._thread: Optional[threading.Thread] = None

    def request(self, pptx_path: Path, strip: bool = False) -> None:
        with self._cond:
            if self._stop.is_set():
                return
            self._pending[(pptx_path, strip)] = None
            self._pending.move_to_end((pptx_path, strip))
            if self._thread is None:
                self._thread = threading.Thread(target=self._run, name="slidedesk-render", daemon=True)
                self._thread.start()
            self._cond.notify()

    def _run(self) -> None:
        while True:
            with self._cond:
                while not self._pending and not self._stop.is_set():
                    self._cond.wait()
                if self._stop.is_set():
                    return
                (pptx_path, strip), _ = self._pending.popitem(last=True)
            try:
                render_deck(pptx_path, self.cache_dir, strip, stop_event=self._stop)
            except (OSError, RenderError) as exc:
                logger.warning("Background rendering of %s failed: %s", pptx_path.name, exc)

    def stop(self, timeout: float = 5.0) -> None:
        with self._cond:
            self._stop.set()
            self._cond.notify_all()
            thread = self._thread
        if thread is not None:
            thread.join(timeout)


def delete_cached(pptx_path: Path, slide_count: int, cache_dir: Path) -> None:
    """Remove cached PNGs of every slide of the current version of `pptx_path`."""
    for strip in (False, True):
        key = _source_key(pptx_path, strip)
        _failures.pop(key, None)
        for index in range(slide_count):
            _cache_file(key, index, cache_dir).unlink(missing_ok=True)


def clear_cache() -> None:
    _failures.clear()
