"""Produce a copy of a .pptx with all slides visible and the layout/theme styling removed."""
from __future__ import annotations

import hashlib
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE

from .pptx_tools import set_slide_hidden

WHITE = RGBColor(255, 255, 255)
BLACK = RGBColor(0, 0, 0)
DARK_FILL_LUMA_THRESHOLD = 48
BRIGHT_BORDER_LUMA_THRESHOLD = 200


def _remove_shape(shape) -> None:
    element = shape.element
    element.getparent().remove(element)


def _shape_has_image(shape) -> bool:
    try:
        return shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    except NotImplementedError:
        return False


def _set_background_white(container) -> None:
    fill = container.background.fill
    fill.solid()
    fill.fore_color.rgb = WHITE


def _luma(rgb: RGBColor) -> float:
    return 0.2126 * rgb[0] + 0.7152 * rgb[1] + 0.0722 * rgb[2]


def _set_fill_white_if_dark(fill, force_if_unknown: bool = False) -> None:
    if fill is None:
        return

    # Some fill proxies (e.g. _NoneFill) raise when foreground color is accessed.
    if getattr(fill, "type", None) is None:
        if force_if_unknown:
            fill.solid()
            fill.fore_color.rgb = WHITE
        return

    try:
        rgb = fill.fore_color.rgb
    except (AttributeError, TypeError, ValueError):
        rgb = None

    if rgb is None:
        if force_if_unknown:
            fill.solid()
            fill.fore_color.rgb = WHITE
        return

    if _luma(rgb) <= DARK_FILL_LUMA_THRESHOLD:
        fill.solid()
        fill.fore_color.rgb = WHITE


def _line_is_explicitly_missing(line) -> bool:
    ln = getattr(line, "_ln", None)
    if ln is None:
        return True
    try:
        return ln.find(".//a:noFill", namespaces=ln.nsmap) is not None
    except (AttributeError, TypeError, ValueError):
        return False


def _set_line_black_if_bright(line, force_if_unknown: bool = False) -> None:
    if line is None or _line_is_explicitly_missing(line):
        return

    fill = getattr(line, "fill", None)
    if fill is None:
        return

    if getattr(fill, "type", None) is None:
        if force_if_unknown:
            fill.solid()
            fill.fore_color.rgb = BLACK
        return

    try:
        rgb = fill.fore_color.rgb
    except (AttributeError, TypeError, ValueError):
        rgb = None

    if rgb is None:
        if force_if_unknown:
            fill.solid()
            fill.fore_color.rgb = BLACK
        return

    if _luma(rgb) >= BRIGHT_BORDER_LUMA_THRESHOLD:
        fill.solid()
        fill.fore_color.rgb = BLACK


def _set_text_frame_black(text_frame) -> None:
    for paragraph in text_frame.paragraphs:
        paragraph.font.color.rgb = BLACK
        for run in paragraph.runs:
            run.font.color.rgb = BLACK


def _set_shape_text_black(shape) -> None:
    has_text_frame = getattr(shape, "has_text_frame", False)
    has_table = getattr(shape, "has_table", False)
    force_unknown_colors = not _shape_has_image(shape)
    shape_fill = getattr(shape, "fill", None)
    has_unknown_fill = getattr(shape_fill, "type", None) is None

    _set_line_black_if_bright(
        getattr(shape, "line", None),
        force_if_unknown=force_unknown_colors and has_unknown_fill,
    )
    _set_fill_white_if_dark(
        shape_fill, force_if_unknown=force_unknown_colors or has_text_frame or has_table
    )

    if has_text_frame:
        _set_text_frame_black(shape.text_frame)

    if has_table:
        for row in shape.table.rows:
            for cell in row.cells:
                _set_fill_white_if_dark(cell.fill, force_if_unknown=True)
                _set_text_frame_black(cell.text_frame)

    try:
        child_shapes = shape.shapes
    except AttributeError:
        child_shapes = None
    if child_shapes is not None:
        for child_shape in child_shapes:
            _set_shape_text_black(child_shape)


def _strip_template_shapes(template_shapes) -> None:
    for shape in list(template_shapes):
        if _shape_has_image(shape) or not getattr(shape, "is_placeholder", False):
            _remove_shape(shape)


def _picture_hashes(shapes) -> set:
    hashes = set()
    for shape in shapes:
        if _shape_has_image(shape):
            hashes.add(hashlib.sha1(shape.image.blob).hexdigest())
        elif getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
            hashes |= _picture_hashes(shape.shapes)
    return hashes


def _remove_pictures(shapes, hashes: set) -> None:
    for shape in list(shapes):
        if _shape_has_image(shape):
            if hashlib.sha1(shape.image.blob).hexdigest() in hashes:
                _remove_shape(shape)
        elif getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
            _remove_pictures(shape.shapes, hashes)


def make_strip_layout_copy(pptx_path: Path, out_path: Path) -> None:
    """Save a copy of `pptx_path` with all slides visible and master/layout decoration removed."""
    presentation = Presentation(str(pptx_path))

    master_images = set()
    for master in presentation.slide_masters:
        master_images |= _picture_hashes(master.shapes)

    containers = []
    for master in presentation.slide_masters:
        containers.append(master)
        containers.extend(master.slide_layouts)

    for container in containers:
        _set_background_white(container)
        _strip_template_shapes(container.shapes)
        for shape in container.shapes:
            _set_shape_text_black(shape)

    for slide in presentation.slides:
        set_slide_hidden(slide, False)
        _remove_pictures(slide.shapes, master_images)
        _set_background_white(slide)
        for shape in slide.shapes:
            _set_shape_text_black(shape)

    presentation.save(str(out_path))
