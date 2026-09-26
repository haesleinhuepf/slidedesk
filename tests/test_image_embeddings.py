import threading
from unittest.mock import Mock

import pytest
from PIL import Image

from slidedesk import SlideProject, embeddings, image_embeddings
from slidedesk.server.app import create_app


@pytest.fixture
def project(tmp_path, monkeypatch):
    project = SlideProject(tmp_path)
    for name in ("deck.pptx", "deck.pdf", "deck.hidden.pdf"):
        (tmp_path / name).touch()
    deck = project.conn.execute(
        "INSERT INTO decks (pptx_path, pptx_mtime, pdf_path, hidden_pdf_path) "
        "VALUES ('deck.pptx', 0, 'deck.pdf', 'deck.hidden.pdf')"
    ).lastrowid
    for i in range(5):
        project.conn.execute(
            "INSERT INTO slides (deck_id, index_in_deck, text, hidden, visible_pdf_page) "
            "VALUES (?, ?, ?, ?, ?)", (deck, i, "" if i == 0 else "text", i == 4, None if i == 4 else i + 1),
        )
    project.conn.commit()
    monkeypatch.setattr(project, "slide_image", lambda sid: Image.new("RGB", (32, 32)))
    monkeypatch.setattr(image_embeddings, "embed_image", Mock(return_value=[1.0, 0.0]))
    try:
        yield project
    finally:
        project.close()


def test_cache_persists_and_invalidates(project):
    assert project.embed_images_pending(limit=2) == 2
    assert project.embed_images_pending() == 3
    assert project.embed_images_pending() == 0
    assert image_embeddings.embed_image.call_count == 5
    reopened = SlideProject(project.folder)
    try:
        assert reopened.embed_images_pending() == 0
    finally:
        reopened.close()
    (project.folder / "deck.hidden.pdf").write_bytes(b"changed hidden export")
    assert project.embed_images_pending() == 1
    (project.folder / "deck.pdf").write_bytes(b"changed visible export")
    assert project.embed_images_pending() == 4
    project.IMAGE_EMBEDDING_MODEL = "new-model"
    assert project.embed_images_pending() == 5
    project.conn.execute("DELETE FROM slides")
    project.conn.commit()
    assert project.conn.execute("SELECT COUNT(*) FROM slide_image_embeddings").fetchone()[0] == 0


@pytest.mark.parametrize("text_available", [False, True])
def test_status_includes_current_vision_coverage(project, monkeypatch, text_available):
    monkeypatch.setattr(project, "_embeddings_available", lambda: text_available)
    monkeypatch.setattr(image_embeddings, "is_available", lambda: True)
    client = create_app(project).test_client()
    status = client.get("/api/embeddings/status").get_json()
    assert status["available"] is text_available
    assert status["total"] == 4
    assert status["vision"] == {"available": True, "total": 5, "embedded": 0}
    image_embeddings.embed_image.assert_not_called()
    project.embed_images_pending(limit=2)
    assert client.get("/api/embeddings/status").get_json()["vision"]["embedded"] == 2
    project.embed_images_pending()
    assert project.embedding_status()["vision"]["embedded"] == 5
    (project.folder / "deck.hidden.pdf").write_bytes(b"changed")
    assert project.embedding_status()["vision"]["embedded"] == 4
    project.IMAGE_EMBEDDING_MODEL = "different-model"
    assert project.embedding_status()["vision"]["embedded"] == 0


def test_unrenderable_slides_do_not_starve_others(project, monkeypatch):
    def render(sid):
        if sid < 3:
            raise FileNotFoundError("PDF unavailable")
        return Image.new("RGB", (32, 32))
    monkeypatch.setattr(project, "slide_image", render)
    assert project.embed_images_pending(limit=1) == 1
    assert project.embed_images_pending() == 2


@pytest.mark.parametrize("change", ["delete", "pdf"])
def test_concurrent_refresh_discards_vector(project, monkeypatch, change):
    def embed(image, model):
        if change == "delete":
            with project._conn_lock:
                project.conn.execute("DELETE FROM slides")
                project.conn.commit()
        else:
            with (project.folder / "deck.pdf").open("ab") as pdf:
                pdf.write(b"changed")
            with (project.folder / "deck.hidden.pdf").open("ab") as pdf:
                pdf.write(b"changed")
        return [1.0, 0.0]
    monkeypatch.setattr(image_embeddings, "embed_image", embed)
    assert project.embed_images_pending() == 0


def seed_vectors(project):
    project.embed_images_pending()
    for sid, vector in enumerate(([1, 0], [1, 0], [0.9, 0.1], [0, 1], [-1, 0]), 1):
        project.conn.execute(
            "UPDATE slide_image_embeddings SET vector=? WHERE slide_id=?",
            (embeddings.pack_vector(vector), sid),
        )
    for sid, vector in enumerate(([1, 0, 0], [0, 1, 0], [-1, 0, 0], [0.9, 0.1, 0], [1, 0, 0]), 1):
        project.conn.execute(
            "INSERT INTO slide_embeddings VALUES (?, ?, ?, 0)",
            (sid, project.EMBEDDING_MODEL, embeddings.pack_vector(vector)),
        )
    project.conn.commit()


def test_combined_ranking_and_unchanged_api(project):
    seed_vectors(project)
    assert [s.id for s in project.similar_slides(1, top_k=4)] == [2, 3, 5, 4]
    assert project.similar_slides(1, top_k=0) == []
    assert project.similar_slides(999) == []
    response = create_app(project).test_client().get("/api/slides/1/similar")
    assert response.status_code == 200
    results = response.get_json()
    assert len({s["id"] for s in results}) == 4
    assert all(set(s) == {"id", "deck_id", "index_in_deck", "visible_pdf_page", "hidden", "text"}
               for s in results)


@pytest.mark.parametrize("missing", ["image", "text"])
def test_single_cache_fallback(project, missing):
    seed_vectors(project)
    table = "slide_image_embeddings" if missing == "image" else "slide_embeddings"
    project.conn.execute(f"DELETE FROM {table}")
    project.conn.commit()
    expected = [5, 4, 2, 3] if missing == "image" else [2, 3, 4, 5]
    assert [s.id for s in project.similar_slides(1, top_k=4)] == expected


def test_stale_image_vectors_are_not_returned(project):
    seed_vectors(project)
    (project.folder / "deck.hidden.pdf").write_bytes(b"updated")
    project.conn.execute("DELETE FROM slide_embeddings")
    project.conn.commit()
    assert [s.id for s in project.similar_slides(1)] == [2, 3, 4]
    (project.folder / "deck.pdf").write_bytes(b"updated")
    assert project.similar_slides(1) == []


def test_scan_keeps_decks_and_removes_their_cached_vectors(tmp_path, monkeypatch):
    from pptx import Presentation
    from slidedesk import scanner
    monkeypatch.setattr(scanner, "convert_pptx_to_pdf", lambda source, dest: dest.touch())
    monkeypatch.setattr(image_embeddings, "embed_image", Mock(return_value=[1, 0]))
    for name in ("one.pptx", "two.pptx"):
        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[6])
        prs.save(tmp_path / name)
    project = SlideProject(tmp_path)
    monkeypatch.setattr(project, "slide_image", lambda sid: Image.new("RGB", (32, 32)))
    try:
        project.scan()
        assert len(project.decks()) == 2
        assert project.embed_images_pending() == 2
        project.scan()
        assert project.embed_images_pending() == 0
        project.refresh_deck(project.decks()[0].id)
        assert project.embed_images_pending() == 1
        (tmp_path / "one.pptx").unlink()
        project.scan()
        assert len(project.decks()) == 1
        assert project.conn.execute("SELECT COUNT(*) FROM slide_image_embeddings").fetchone()[0] == 1
    finally:
        project.close()


@pytest.mark.parametrize("structured_output", [False, True])
def test_clip_loads_once_and_uses_inference(monkeypatch, structured_output):
    import sys
    from contextlib import nullcontext
    from types import SimpleNamespace
    vector = Mock(spec=["squeeze", "cpu", "tolist"])
    vector.squeeze.return_value.cpu.return_value.tolist.return_value = [1.0, 2.0]
    encoder = Mock()
    encoder.get_image_features.return_value = (
        SimpleNamespace(pooler_output=vector) if structured_output else vector
    )
    processor = Mock(return_value={"pixel_values": "pixels"})
    model_class = Mock()
    model_class.from_pretrained.return_value = encoder
    processor_class = Mock()
    processor_class.from_pretrained.return_value = processor
    no_grad = Mock(side_effect=nullcontext)
    monkeypatch.setitem(sys.modules, "torch", SimpleNamespace(no_grad=no_grad))
    monkeypatch.setitem(sys.modules, "transformers", SimpleNamespace(
        CLIPModel=model_class, CLIPProcessor=processor_class))
    monkeypatch.setattr(image_embeddings, "_models", {})
    with Image.new("RGBA", (32, 32)) as image:
        assert image_embeddings.embed_image(image) == [1, 2]
        assert image_embeddings.embed_image(image) == [1, 2]
    model_class.from_pretrained.assert_called_once_with(image_embeddings.MODEL, low_cpu_mem_usage=False)
    encoder.eval.assert_called_once()
    assert processor.call_args.kwargs["images"].mode == "RGB"
    assert no_grad.call_count == 2
    encoder.get_image_features.assert_called_with(pixel_values="pixels")


def test_real_clip_inference_without_model_download(monkeypatch):
    import math
    pytest.importorskip("torch")
    transformers = pytest.importorskip("transformers")
    config = transformers.CLIPConfig(
        projection_dim=16,
        text_config={"hidden_size": 32, "intermediate_size": 64,
                     "num_hidden_layers": 1, "num_attention_heads": 4, "vocab_size": 100},
        vision_config={"hidden_size": 32, "intermediate_size": 64,
                       "num_hidden_layers": 1, "num_attention_heads": 4,
                       "image_size": 32, "patch_size": 16},
    )
    model = transformers.CLIPModel(config)
    processor = transformers.CLIPImageProcessor(
        size={"shortest_edge": 32}, crop_size={"height": 32, "width": 32})
    monkeypatch.setattr(transformers.CLIPModel, "from_pretrained", lambda name, **kwargs: model)
    monkeypatch.setattr(transformers.CLIPProcessor, "from_pretrained", lambda name: processor)
    monkeypatch.setattr(image_embeddings, "_models", {})
    with Image.new("RGB", (64, 32), "red") as image:
        vector = image_embeddings.embed_image(image)
        assert image_embeddings.embed_image(image) == vector
    assert len(vector) == 16
    assert all(math.isfinite(value) for value in vector)
    assert not model.training
