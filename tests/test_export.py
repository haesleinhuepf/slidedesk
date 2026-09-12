from zipfile import ZipFile

from pptx import Presentation

from slidedesk import SlideProject
from slidedesk.pptx_tools import extract_text, is_slide_hidden, set_slide_hidden
from slidedesk.scanner import _index_slides
from slidedesk.server.app import create_app


def test_export_selected_slide_ids_in_order(tmp_path):
    project = SlideProject(tmp_path)
    try:
        decks = []
        for name in ("first", "second"):
            path = tmp_path / f"{name}.pptx"
            prs = Presentation()
            for index in range(4):
                slide = prs.slides.add_slide(prs.slide_layouts[0])
                slide.shapes.title.text = f"{name} slide {index + 1}"
                set_slide_hidden(slide, index == 1)
            prs.save(path)
            deck_id = project.conn.execute(
                "INSERT INTO decks (pptx_path, pptx_mtime) VALUES (?, ?)",
                (path.name, path.stat().st_mtime),
            ).lastrowid
            _index_slides(project.conn, deck_id, path)
            decks.append(project.slides(deck_id))
        project.conn.commit()
        # Nonconsecutive positions, a hidden slide, both decks, and a repeat.
        selected = [decks[0][3], decks[1][1], decks[0][0], decks[1][2], decks[0][3]]
        client = create_app(project).test_client()
        response = client.post("/api/export", json={
            "slide_ids": [slide.id for slide in selected]
        })
        assert response.status_code == 200
        assert response.json["path"] == "export_1.pptx"
        output = project.export_dir / response.json["path"]
        assert not output.is_relative_to(project.folder)
        assert not (tmp_path / "export_1.pptx").exists()
        download = client.get(f"/api/export/{response.json['path']}/download")
        assert download.status_code == 200
        assert download.data == output.read_bytes()
        assert 'filename=export_1.pptx' in download.headers['Content-Disposition']
        download.close()
        assert client.get("/api/export/first.pptx/download").status_code == 404
        with ZipFile(output) as archive:
            names = archive.namelist()
            assert len(names) == len(set(names)), "Duplicate PPTX package parts"
        exported = Presentation(output)
        assert [extract_text(slide) for slide in exported.slides] == [
            slide.text for slide in selected
        ]
        assert all(not is_slide_hidden(slide) for slide in exported.slides)
        original_bytes = output.read_bytes()
        second = client.post("/api/export", json={"slide_ids": [selected[0].id]})
        assert second.status_code == 200
        assert second.json["path"] == "export_2.pptx"
        assert (project.export_dir / "export_2.pptx").exists()
        assert output.read_bytes() == original_bytes
        assert sorted(path.name for path in tmp_path.glob("*.pptx")) == [
            "first.pptx", "second.pptx"
        ]
    finally:
        project.close()
    assert not project.export_dir.exists()
