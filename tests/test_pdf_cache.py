from pathlib import Path

from pptx import Presentation

from slidedesk import SlideProject, scanner
from slidedesk.convert import ConversionError


def test_pdf_exports_mirror_source_tree_and_refresh(tmp_path, monkeypatch):
    calls = []

    def convert(source, destination):
        calls.append(destination)
        destination.write_bytes(b"pdf")

    monkeypatch.setattr(scanner, "convert_pptx_to_pdf", convert)
    for folder in ("", "talks/nested", ".slidedesk/_cache", ".slidedesk_cache"):
        source = tmp_path / folder / "deck.pptx"
        source.parent.mkdir(parents=True, exist_ok=True)
        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[1])
        prs.save(source)

    project = SlideProject(tmp_path)
    try:
        project.scan()
        assert len(project.decks()) == 2
        assert len(calls) == 6
        for deck in project.decks():
            expected = Path(".slidedesk/_cache") / deck.pptx_path
            assert deck.pdf_path == expected.with_suffix(".pdf").as_posix()
            assert deck.hidden_pdf_path == expected.with_suffix(".hidden.pdf").as_posix()
            assert deck.strip_pdf_path == expected.with_suffix(".strip-layout.hidden.pdf").as_posix()
            assert (tmp_path / deck.pdf_path).exists()
            assert (tmp_path / deck.hidden_pdf_path).exists()
            assert (tmp_path / deck.strip_pdf_path).exists()
            assert not (tmp_path / deck.pptx_path).with_suffix(".pdf").exists()
        project.scan()
        assert len(calls) == 6
        project.refresh_deck(project.decks()[0].id)
        assert len(calls) == 9
    finally:
        project.close()


def test_scan_relocates_indexed_sidecars_without_conversion(tmp_path, monkeypatch):
    source = tmp_path / "talks" / "deck.pptx"
    source.parent.mkdir()
    Presentation().save(source)
    project = SlideProject(tmp_path)
    try:
        for suffix in (".pdf", ".hidden.pdf"):
            source.with_suffix(suffix).write_bytes(suffix.encode())
        project.conn.execute(
            "INSERT INTO decks (pptx_path, pptx_mtime, pdf_path, hidden_pdf_path) "
            "VALUES (?, ?, ?, ?)",
            ("talks/deck.pptx", source.stat().st_mtime,
             "talks/deck.pdf", "talks/deck.hidden.pdf"),
        )
        project.conn.commit()
        strip = tmp_path / ".slidedesk/_cache/talks/deck.strip-layout.hidden.pdf"
        strip.parent.mkdir(parents=True, exist_ok=True)
        strip.write_bytes(b"strip")

        def unexpected_conversion(*args):
            raise AssertionError("Existing exports should be reused")

        monkeypatch.setattr(scanner, "convert_pptx_to_pdf", unexpected_conversion)
        project.scan()
        deck = project.decks()[0]
        for relative, suffix in ((deck.pdf_path, ".pdf"),
                                 (deck.hidden_pdf_path, ".hidden.pdf")):
            assert (tmp_path / relative).read_bytes() == suffix.encode()
            assert not source.with_suffix(suffix).exists()
    finally:
        project.close()


def test_failed_conversion_is_persisted_and_not_retried(tmp_path, monkeypatch):
    source = tmp_path / "deck.pptx"
    Presentation().save(source)
    calls = []

    def fail_conversion(source_path, destination):
        calls.append(source_path)
        raise ConversionError("conversion failed")

    monkeypatch.setattr(scanner, "convert_pptx_to_pdf", fail_conversion)
    project = SlideProject(tmp_path)
    try:
        project.scan()
        assert len(calls) == 1
        row = project.conn.execute(
            "SELECT pdf_conversion_failed FROM decks WHERE pptx_path = ?",
            ("deck.pptx",),
        ).fetchone()
        assert row["pdf_conversion_failed"] == 1

        project.scan()
        assert len(calls) == 1
    finally:
        project.close()


def test_strip_conversion_retries_after_delay(tmp_path, monkeypatch):
    source = tmp_path / "deck.pptx"
    Presentation().save(source)
    strip_attempts = []
    sleeps = []

    def convert(source_path, destination):
        if destination.name.endswith(".strip-layout.hidden.pdf"):
            strip_attempts.append(source_path)
            if len(strip_attempts) == 1:
                raise ConversionError("temporary conversion failure")
        destination.write_bytes(b"pdf")

    monkeypatch.setattr(scanner, "convert_pptx_to_pdf", convert)
    monkeypatch.setattr(scanner.time, "sleep", sleeps.append)
    project = SlideProject(tmp_path)
    try:
        project.scan()

        assert len(strip_attempts) == 2
        assert sleeps == [0.5]
        row = project.conn.execute(
            "SELECT pdf_conversion_failed FROM decks WHERE pptx_path = ?",
            ("deck.pptx",),
        ).fetchone()
        assert row["pdf_conversion_failed"] == 0
        assert (project.cache_dir / "deck.strip-layout.hidden.pdf").exists()
    finally:
        project.close()
