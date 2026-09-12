import pytest
from pptx import Presentation

from slidedeck import SlideProject
from slidedeck.scanner import _index_slides
from slidedeck.server.app import create_app


@pytest.fixture
def project(tmp_path):
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    prs.slides.add_slide(prs.slide_layouts[6])._element.set("show", "0")
    prs.save(tmp_path / "deck.pptx")
    project = SlideProject(tmp_path)
    project.conn.execute(
        "INSERT INTO decks (pptx_path, pptx_mtime, pdf_path, hidden_pdf_path) "
        "VALUES ('deck.pptx', 0, 'deck.pdf', 'deck.hidden.pdf')"
    )
    _index_slides(project.conn, 1, tmp_path / "deck.pptx")
    project.conn.commit()
    yield project
    project.close()


def test_visibility_actions_and_persistence(project):
    client = create_app(project).test_client()
    first, second = project.slides(1)
    assert client.patch(f"/api/slides/{first.id}/hidden", json={"hidden": True}).status_code == 200
    assert project.slides(1, include_hidden=False) == []
    assert client.patch("/api/decks/1/slides/hidden", json={"hidden": False}).status_code == 200
    assert len(project.slides(1, include_hidden=False)) == 2
    assert client.patch("/api/decks/1/slides/hidden", json={"hidden": True}).status_code == 200
    assert all(s.hidden for s in project.slides(1))
    client.patch(f"/api/slides/{second.id}/hidden", json={"hidden": False})
    assert [s.hidden for s in project.slides(1)] == [True, False]
    assert "hidden" not in client.get("/api/decks").get_json()[0]
    # Visibility follows the source slide identity even when slides are reordered.
    prs = Presentation(project.folder / "deck.pptx")
    ids = prs.slides._sldIdLst
    ids.insert(0, ids[-1])
    prs.save(project.folder / "deck.pptx")
    _index_slides(project.conn, 1, project.folder / "deck.pptx")
    project.conn.commit()
    assert [s.hidden for s in project.slides(1)] == [False, True]
    reopened = SlideProject(project.folder)
    try:
        assert [s.hidden for s in reopened.slides(1)] == [False, True]
    finally:
        reopened.close()


@pytest.mark.parametrize("url", ["/api/slides/1/hidden", "/api/decks/1/slides/hidden"])
@pytest.mark.parametrize("body", [{}, {"hidden": "false"}, {"hidden": 0}, [], None])
def test_invalid_visibility(project, url, body):
    assert create_app(project).test_client().patch(url, json=body).status_code == 400
    assert [s.hidden for s in project.slides(1)] == [False, True]


@pytest.mark.parametrize("url", ["/api/slides/999/hidden", "/api/decks/999/slides/hidden"])
def test_missing_visibility_target(project, url):
    assert create_app(project).test_client().patch(url, json={"hidden": True}).status_code == 404


def test_visibility_does_not_change_render_source(project, monkeypatch):
    from slidedeck import images

    monkeypatch.setattr(images, "render_page", lambda path, page, cache, dpi: (path.name, page))
    first, second = project.slides(1)
    project.set_slide_hidden(first.id, True)
    project.set_slide_hidden(second.id, False)
    assert project.slide_image(first.id) == ("deck.pdf", 1)
    assert project.slide_image(second.id) == ("deck.hidden.pdf", 2)
