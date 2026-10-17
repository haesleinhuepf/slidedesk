import sys
import threading
from types import SimpleNamespace

import pytest
from PIL import Image
from pptx import Presentation

from slidedesk import images


@pytest.fixture
def deck(tmp_path):
    path = tmp_path / "deck.pptx"
    prs = Presentation()
    for _ in range(2):
        prs.slides.add_slide(prs.slide_layouts[6])
    prs.slides[1]._element.set("show", "0")
    prs.save(path)
    images.clear_cache()
    yield path
    images.clear_cache()


def fake_render(monkeypatch, calls):
    def render_slides(prs, indexes):
        calls.append([(i, prs.slides[i]._element.get("show")) for i in indexes])
        return [Image.new("RGB", (4, 4), "red") for _ in indexes]

    monkeypatch.setitem(sys.modules, "render_pptx", SimpleNamespace(render_slides=render_slides))


def test_render_slide_image_renders_hidden_slides_and_caches_pngs(deck, monkeypatch):
    calls = []
    fake_render(monkeypatch, calls)
    cache = deck.parent / "cache"

    first = images.render_slide_image(deck, 0, cache)
    second = images.render_slide_image(deck, 1, cache)

    assert first.size == second.size == (4, 4)
    assert first is not second
    assert calls == [[(0, None)], [(1, None)]]
    assert len(list(cache.glob("*.png"))) == 2
    assert not list(deck.parent.glob("*.pdf"))

    images.clear_cache()
    assert images.render_slide_image(deck, 0, cache).size == (4, 4)
    assert len(calls) == 2


def test_render_slide_image_strip_is_grayscale(deck, monkeypatch):
    fake_render(monkeypatch, [])
    pixel = images.render_slide_image(deck, 0, deck.parent / "cache", strip=True).getpixel((0, 0))
    assert pixel[0] == pixel[1] == pixel[2]
    assert images.cached_path(deck, 0, deck.parent / "cache", strip=False) is None


def test_render_slide_image_rejects_missing_slide(deck, monkeypatch):
    fake_render(monkeypatch, [])
    with pytest.raises(ValueError):
        images.render_slide_image(deck, 5, deck.parent / "cache")


def test_render_deck_skips_cached_slides_and_honours_stop(deck, monkeypatch):
    calls = []
    fake_render(monkeypatch, calls)
    cache = deck.parent / "cache"

    stop = threading.Event()
    stop.set()
    images.render_deck(deck, cache, stop_event=stop)
    assert calls == []

    images.render_slide_image(deck, 0, cache)
    images.render_deck(deck, cache)
    assert calls == [[(0, None)], [(1, None)]]


def test_render_failure_is_remembered_until_cleared(deck, monkeypatch):
    def failing_render(prs, indexes):
        raise RuntimeError("boom")

    monkeypatch.setitem(sys.modules, "render_pptx", SimpleNamespace(render_slides=failing_render))
    cache = deck.parent / "cache"
    with pytest.raises(images.RenderError):
        images.render_deck(deck, cache)
    assert images.render_failure(deck) == "RuntimeError: boom"
    assert images.render_failure(deck, strip=True) is None

    fake_render(monkeypatch, [])
    with pytest.raises(images.RenderError, match="failed earlier"):
        images.render_deck(deck, cache)
    images.delete_cached(deck, 2, cache)
    assert images.render_failure(deck) is None
    images.render_deck(deck, cache)
    assert images.cached_path(deck, 1, cache) is not None


def test_background_renderer_renders_requested_deck(deck, monkeypatch):
    calls = []
    fake_render(monkeypatch, calls)
    cache = deck.parent / "cache"
    renderer = images.BackgroundRenderer(cache)
    try:
        renderer.request(deck)
        for _ in range(200):
            if images.cached_path(deck, 1, cache) is not None:
                break
            threading.Event().wait(0.01)
        assert images.cached_path(deck, 0, cache) is not None
        assert images.cached_path(deck, 1, cache) is not None
    finally:
        renderer.stop()
    renderer.request(deck, strip=True)  # ignored after stop
    assert images.cached_path(deck, 0, cache, strip=True) is None
