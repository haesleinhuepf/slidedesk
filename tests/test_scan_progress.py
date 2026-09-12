from concurrent.futures import ThreadPoolExecutor
import threading

from pptx import Presentation

from slidedeck import SlideProject, scanner
from slidedeck.server.app import create_app


def test_decks_available_while_next_deck_is_converting(tmp_path, monkeypatch):
    paths = [tmp_path / name for name in ('first.pptx', 'second.pptx')]
    for path in paths:
        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[1])
        prs.save(path)
    monkeypatch.setattr(scanner, '_iter_pptx_files', lambda root: iter(paths))
    converting = threading.Event()
    release = threading.Event()

    def convert(source, destination):
        if source == paths[1]:
            converting.set()
            assert release.wait(10)
        destination.touch()

    monkeypatch.setattr(scanner, 'convert_pptx_to_pdf', convert)
    project = SlideProject(tmp_path)
    app = create_app(project)
    try:
        with ThreadPoolExecutor(max_workers=2) as executor:
            scan = executor.submit(project.scan)
            try:
                assert converting.wait(5)
                def read_decks():
                    with app.test_client() as client:
                        return client.get('/api/decks').get_json()
                decks = executor.submit(read_decks).result(timeout=2)
                assert [deck['name'] for deck in decks] == ['first.pptx']
                assert len(decks[0]['slides']) == 1
            finally:
                release.set()
            scan.result(timeout=5)
        assert len(project.decks()) == 2
    finally:
        project.close()
